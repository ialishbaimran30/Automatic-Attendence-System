from datetime import date
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from .models import AttendanceSession, SessionRoster, AttendanceRecord

def date_range(query):
    try:
        start = date.fromisoformat(query['start']) if query.get('start') else None
        end = date.fromisoformat(query['end']) if query.get('end') else None
        if start and end and end < start: raise ValueError()
        month = query.get('month')
        if month:
            year, m = map(int, month.split('-'))
            import calendar
            start, end = date(year, m, 1), date(year, m, calendar.monthrange(year, m)[1])
    except (ValueError, KeyError): raise ValidationError('Use ISO start/end dates or month=YYYY-MM.')
    return start, end

def grid(section, query):
    start, end = date_range(query)
    sessions = AttendanceSession.objects.filter(section=section, closed_at__isnull=False).order_by('scheduled_start')
    if start: sessions = sessions.filter(scheduled_start__date__gte=start)
    if end: sessions = sessions.filter(scheduled_start__date__lte=end)
    if query.get('semester') and query['semester'] != section.semester: sessions = sessions.none()
    sessions = list(sessions)
    ids = [s.pk for s in sessions]
    snapshots = SessionRoster.objects.filter(session_id__in=ids).order_by('session__scheduled_start')
    students, eligible = {}, {}
    for row in snapshots:
        students[row.student_id] = {'student': row.student_id, 'name': row.name, 'registration_number': row.registration_number}
        eligible.setdefault(row.student_id, set()).add(str(row.session_id))
    # Show current students even when no closed session exists (percentage is null, not 0).
    for student in section.students.all():
        students.setdefault(student.pk, {'student': student.pk, 'name': student.name, 'registration_number': student.registration_number})
    records = {(r.student_id, str(r.session_id)): r.status for r in AttendanceRecord.objects.filter(session_id__in=ids)}
    rows = []
    for student_id, info in students.items():
        denominator = len(eligible.get(student_id, set()))
        statuses = {str(s.pk): records.get((student_id, str(s.pk))) if str(s.pk) in eligible.get(student_id, set()) else None for s in sessions}
        attended = sum(s in ('present', 'late') for s in statuses.values())
        percentage = round(100 * attended / denominator, 2) if denominator else None
        rows.append({**info, 'statuses': statuses, 'attended': attended, 'denominator': denominator,
                     'percentage': percentage, 'below_threshold': percentage is not None and percentage < section.setting('attendance_threshold')})
    return {'section': section.pk, 'semester': section.semester, 'threshold': section.setting('attendance_threshold'),
            'late_treatment': 'Late counts as attended; only closed sessions whose snapshots include the student form the denominator.',
            'sessions': [{'id': str(s.pk), 'date': timezone.localtime(s.scheduled_start).isoformat()} for s in sessions],
            'students': sorted(rows, key=lambda s: s['registration_number'])}

def workbook(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = 'Attendance'
    sheet.append(['Registration', 'Student', *[s['date'] for s in data['sessions']], 'Attended', 'Eligible sessions', 'Percentage', 'Below threshold'])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17324D')
    def safe(value):
        # Prevent formula injection from registration/name strings.
        return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
    for row in data['students']:
        sheet.append([safe(row['registration_number']), safe(row['name']), *[row['statuses'][s['id']] or 'N/A' for s in data['sessions']],
                      row['attended'], row['denominator'], row['percentage'], row['below_threshold']])
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for idx in range(1, sheet.max_column + 1): sheet.column_dimensions[get_column_letter(idx)].width = 22
    notes = book.create_sheet('Definitions')
    notes.append(['Threshold (%)', data['threshold']])
    notes.append(['Semester', data['semester']])
    notes.append(['Calculation', data['late_treatment']])
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 110
    return book
