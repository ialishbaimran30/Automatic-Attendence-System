import io
import tempfile
import uuid
from datetime import timedelta
from unittest.mock import patch
from django.test import TestCase, override_settings
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from openpyxl import load_workbook
from PIL import Image
from attendance.models import *
from attendance import services
from attendance.tasks import recognize, purge_expired
from recognition.engine import RecognitionRejected
from recognition.storage import write_private, private_path, encrypt, encode_embedding

def jpeg():
    stream = io.BytesIO()
    Image.new('RGB', (100, 100), (100, 120, 130)).save(stream, format='JPEG')
    return stream.getvalue()

class AttendanceTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user('teacher', password='password', role='teacher')
        self.other = User.objects.create_user('other', password='password', role='teacher')
        self.admin = User.objects.create_user('admin', password='password', role='admin')
        self.department = Department.objects.create(name='CS')
        self.course = Course.objects.create(department=self.department, code='CS1', name='Vision')
        self.section = Section.objects.create(course=self.course, name='A', semester='Fall')
        self.section.teachers.add(self.teacher)
        self.foreign = Section.objects.create(course=self.course, name='B', semester='Fall')
        self.foreign.teachers.add(self.other)
        self.student = Student.objects.create(name='Ayesha', registration_number='CS01', department=self.department, consent_at=timezone.now())
        self.second = Student.objects.create(name='Bilal', registration_number='CS02', department=self.department)
        self.outsider = Student.objects.create(name='Foreign', registration_number='CS03', department=self.department)
        Membership.objects.create(student=self.student, section=self.section)
        Membership.objects.create(student=self.second, section=self.section)
        Membership.objects.create(student=self.outsider, section=self.foreign)
        now = timezone.localtime().replace(hour=12)
        self.clock = patch('django.utils.timezone.now', return_value=now)
        self.clock.start()
        self.addCleanup(self.clock.stop)
        self.schedule = Schedule.objects.create(section=self.section, weekday=now.weekday(), start_time=(now - timedelta(minutes=1)).time(),
            end_time=(now + timedelta(hours=1)).time(), valid_from=now.date(), valid_until=now.date())
        self.client = APIClient()
        self.client.force_authenticate(self.teacher)
        self.session = services.start_session(self.teacher, self.section, self.schedule)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.storage = override_settings(PRIVATE_STORAGE=self.temp.name)
        self.storage.enable()
        self.addCleanup(self.storage.disable)
    def path(self, action): return f'/api/sessions/{self.session.pk}/{action}/'
    def review(self, **kwargs):
        values = dict(session=self.session, requested_by=self.teacher, status='review', proposed_student=self.student,
                      captured_at=timezone.now(), expires_at=timezone.now() + timedelta(minutes=5), score=0.91)
        values.update(kwargs)
        return RecognitionJob.objects.create(**values)

    def test_students_have_no_login_or_attendance_access(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.student.pk, 'reason': 'x'}).status_code, 403)
        self.assertEqual(self.client.get(f'/api/students/{self.student.pk}/percentages/').status_code, 403)
        csrf_client = APIClient(enforce_csrf_checks=True)
        self.assertEqual(csrf_client.post('/api/auth/login/', {'username': self.student.registration_number, 'password': 'password'}).status_code, 403)
        self.assertFalse(User.objects.filter(username=self.student.registration_number).exists())

    def test_teacher_scoped_and_admin_all(self):
        self.assertEqual(len(self.client.get('/api/sections/').data), 1)
        self.assertEqual(self.client.get(f'/api/sections/{self.foreign.pk}/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/sections/{self.foreign.pk}/export/').status_code, 404)
        self.assertEqual(self.client.post('/api/courses/', {'name': 'x'}).status_code, 403)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.student.pk, 'reason': 'x'}).status_code, 404)
        self.assertEqual(self.client.get(f'/api/students/{self.student.pk}/percentages/').status_code, 404)
        self.client.force_authenticate(self.admin)
        self.assertEqual(len(self.client.get('/api/sections/').data), 2)

    def test_manual_roster_duplicate_reason_and_close(self):
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.outsider.pk, 'reason': 'x'}).status_code, 400)
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.student.pk, 'reason': ''}).status_code, 400)
        response = self.client.post(self.path('manual'), {'student': self.student.pk, 'reason': 'Visual verification', 'timestamp': '2000-01-01'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(AttendanceRecord.objects.get(student=self.student).timestamp, timezone.now())
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.student.pk, 'reason': 'again'}).status_code, 400)
        self.assertEqual(self.client.get(f'/api/sessions/{self.session.pk}/').data['checked_in_count'], 1)
        self.assertEqual(self.client.post(self.path('close'), {}).status_code, 200)
        self.assertEqual(AttendanceRecord.objects.get(student=self.second).status, 'absent')
        self.assertEqual(self.client.post(self.path('manual'), {'student': self.second.pk, 'reason': 'x'}).status_code, 400)
        self.assertEqual(self.client.post(self.path('close'), {}).status_code, 400)

    def test_mandatory_confirmation_capture_time_and_duplicate(self):
        job = self.review(captured_at=self.session.scheduled_start + timedelta(minutes=5))
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        response = self.client.post(self.path('confirm'), {'job': str(job.pk)})
        self.assertEqual(response.status_code, 201)
        record = AttendanceRecord.objects.get(student=self.student)
        self.assertEqual(record.status, 'present')
        self.assertEqual(record.timestamp, job.captured_at)
        self.assertEqual(self.client.post(self.path('confirm'), {'job': str(job.pk)}).status_code, 400)
        second = self.review()
        self.assertEqual(self.client.post(self.path('confirm'), {'job': str(second.pk)}).status_code, 400)

    def test_stale_foreign_rejected_and_wrong_user_jobs(self):
        for changes in [{'expires_at': timezone.now() - timedelta(seconds=1)}, {'status': 'rejected'},
                        {'proposed_student': self.outsider}, {'requested_by': self.other}]:
            job = self.review(**changes)
            self.assertEqual(self.client.post(self.path('confirm'), {'job': str(job.pk)}).status_code, 400)
        foreign_schedule = Schedule.objects.create(section=self.foreign, weekday=self.schedule.weekday,
            start_time=self.schedule.start_time, end_time=self.schedule.end_time, valid_from=self.schedule.valid_from, valid_until=self.schedule.valid_until)
        foreign_session = services.start_session(self.other, self.foreign, foreign_schedule)
        foreign_job = self.review(session=foreign_session)
        self.assertEqual(self.client.post(self.path('confirm'), {'job': str(foreign_job.pk)}).status_code, 400)
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_grace_boundary_and_correction_audited(self):
        boundary = self.session.scheduled_start + timedelta(minutes=self.session.grace_minutes)
        self.assertEqual(services.status_at(self.session, boundary), 'present')
        self.assertEqual(services.status_at(self.session, boundary + timedelta(microseconds=1)), 'late')
        job = self.review(captured_at=boundary + timedelta(seconds=1))
        self.client.post(self.path('confirm'), {'job': str(job.pk)})
        self.assertEqual(AttendanceRecord.objects.get(student=self.student).status, 'late')
        self.client.post(self.path('close'), {})
        self.assertEqual(self.client.post(self.path('correct'), {'student': self.second.pk, 'status': 'present', 'reason': ''}).status_code, 400)
        self.assertEqual(self.client.post(self.path('correct'), {'student': self.second.pk, 'status': 'present', 'reason': 'Verified paper register'}).status_code, 200)
        self.assertTrue(AuditLog.objects.filter(event='attendance.corrected', details__reason='Verified paper register').exists())

    def test_snapshot_percentage_late_and_exports(self):
        self.assertIsNone(self.client.get(f'/api/sections/{self.section.pk}/grid/').data['students'][0]['percentage'])
        self.review(captured_at=self.session.scheduled_start + timedelta(minutes=11))
        services.confirm_job(self.teacher, self.session.pk, RecognitionJob.objects.first().pk)
        Membership.objects.filter(student=self.student).delete()
        self.student.name = '=Formula'
        self.student.save()
        self.client.post(self.path('close'), {})
        response = self.client.get(f'/api/sections/{self.section.pk}/grid/')
        rows = {s['student']: s for s in response.data['students']}
        self.assertEqual(rows[self.student.pk]['percentage'], 100)
        self.assertEqual(rows[self.student.pk]['name'], 'Ayesha')
        self.assertTrue(rows[self.second.pk]['below_threshold'])
        self.assertEqual(rows[self.second.pk]['denominator'], 1)
        xlsx = self.client.get(f'/api/sections/{self.section.pk}/export/')
        self.assertEqual(xlsx.status_code, 200)
        book = load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(book['Attendance'].max_row, 3)
        self.assertEqual(book['Attendance'].cell(2, 6).value, 100)
        self.assertEqual(self.client.get(f'/api/sections/{self.section.pk}/grid/?month=invalid').status_code, 400)

    def test_capture_grants_single_use_bound_expiry_and_queue_failure(self):
        token, _ = services.issue_capture(self.teacher, self.session.pk)
        form = {'token': token, 'frame': SimpleUploadedFile('f.jpg', jpeg(), content_type='image/jpeg')}
        with patch('attendance.tasks.recognize.delay', side_effect=RuntimeError('broker gone')):
            response = self.client.post(self.path('capture'), form, format='multipart')
        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.data['status'], 'error')
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        self.assertFalse(list(private_path('frames').glob('*')))
        self.assertEqual(self.client.post(self.path('capture'), {'token': token, 'frame': SimpleUploadedFile('f.jpg', jpeg())}, format='multipart').status_code, 400)
        wrong_token, _ = services.issue_capture(self.teacher, self.session.pk)
        with self.assertRaises(Exception): services.create_job(self.other, self.session.pk, wrong_token, '')
        CaptureGrant.objects.filter(token_hash=services.token_hash(wrong_token)).update(expires_at=timezone.now() - timedelta(seconds=1))
        with self.assertRaises(Exception): services.create_job(self.teacher, self.session.pk, wrong_token, '')

    def test_worker_fail_closed_and_frame_deleted(self):
        for exc in [RecognitionRejected('liveness_failed'), RecognitionRejected('exactly_one_face_required'),
                    RecognitionRejected('unknown_face'), RecognitionRejected('uncertain_match'), RuntimeError('model failed')]:
            path = write_private(jpeg())
            job = RecognitionJob.objects.create(session=self.session, requested_by=self.teacher, expires_at=timezone.now()+timedelta(minutes=5), frame_path=path)
            with patch('attendance.tasks.extract_embedding', side_effect=exc): recognize(str(job.pk))
            job.refresh_from_db()
            self.assertIn(job.status, ('rejected', 'error'))
            self.assertFalse(private_path(path).exists())
            self.assertEqual(AttendanceRecord.objects.count(), 0)
            self.assertTrue(AuditLog.objects.filter(details__job_id=str(job.pk)).exists())

    def test_worker_success_proposes_only_and_close_race(self):
        path = write_private(jpeg())
        job = RecognitionJob.objects.create(session=self.session, requested_by=self.teacher, expires_at=timezone.now()+timedelta(minutes=5), frame_path=path)
        with patch('attendance.tasks.extract_embedding', return_value=[1]), patch('attendance.tasks.match_embedding', return_value=(self.student.pk, .92)) as matching:
            recognize(str(job.pk))
            self.assertEqual(set(matching.call_args.args[1]), {self.student.pk, self.second.pk})
        job.refresh_from_db()
        self.assertEqual(job.status, 'review')
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        services.close_session(self.teacher, self.session.pk)
        self.assertEqual(self.client.post(self.path('confirm'), {'job': str(job.pk)}).status_code, 400)

    def test_capture_poll_and_confirm_full_api_flow(self):
        token = self.client.post(self.path('capture-grant'), {}).data['token']
        with patch('attendance.tasks.extract_embedding', return_value=[1]), patch('attendance.tasks.match_embedding', return_value=(self.student.pk,.92)):
            response = self.client.post(self.path('capture'), {'token':token,'frame':SimpleUploadedFile('f.jpg',jpeg())}, format='multipart')
        self.assertEqual(response.status_code, 202)
        job_id = response.data['id']
        review = self.client.get(f'/api/jobs/{job_id}/')
        self.assertEqual(review.data['status'], 'review')
        self.assertEqual(review.data['name'], self.student.name)
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        self.assertEqual(self.client.post(self.path('confirm'), {'job':job_id}).status_code, 201)
        self.assertEqual(AttendanceRecord.objects.count(), 1)

    def test_enrollment_single_use_expiry_consent_no_attendance(self):
        self.client.force_authenticate(None)
        link = EnrollmentLink.objects.create(student=self.student, issued_by=self.teacher, token_hash=services.token_hash('valid'), expires_at=timezone.now()+timedelta(minutes=30))
        vector = [1.] + [0.] * 511
        def submit(token, consent='true'):
            return self.client.post('/api/enroll/', {'token': token, 'consent': consent, 'photos': SimpleUploadedFile('f.jpg', jpeg())}, format='multipart')
        with patch('attendance.enrollment.extract_embedding', return_value=__import__('numpy').array(vector)):
            self.assertEqual(submit('valid', 'false').status_code, 400)
            self.assertEqual(submit('valid').status_code, 201)
            self.assertEqual(submit('valid').status_code, 400)
            EnrollmentLink.objects.create(student=self.second, issued_by=self.teacher, token_hash=services.token_hash('expired'), expires_at=timezone.now()-timedelta(seconds=1))
            self.assertEqual(submit('expired').status_code, 400)
        sample = BiometricSample.objects.get()
        self.assertNotIn(jpeg(), bytes(sample.photo_ciphertext))
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        self.assertEqual(self.client.post(self.path('confirm'), {'job': str(uuid.uuid4())}).status_code, 403)
        self.client.force_authenticate(self.teacher)
        data = self.client.get(f'/api/students/{self.student.pk}/').data
        self.assertNotIn('embedding', str(data))
        self.assertEqual(self.client.post(f'/api/students/{self.outsider.pk}/enrollment-link/', {}).status_code, 404)

    def test_retention_purge_and_biometric_deletion(self):
        BiometricSample.objects.create(student=self.student, embedding_ciphertext=encode_embedding([1]+[0]*511),
            photo_ciphertext=encrypt(b'x'), capture_group=uuid.uuid4(), photo_sha256='a'*64, expires_at=timezone.now()-timedelta(days=1))
        purge_expired()
        self.assertEqual(BiometricSample.objects.count(), 0)
        self.assertTrue(AuditLog.objects.filter(event='biometrics.retention_purge').exists())
        self.assertEqual(self.client.post(f'/api/students/{self.student.pk}/purge-biometrics/', {'reason': 'withdrawal'}).status_code, 403)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.post(f'/api/students/{self.student.pk}/purge-biometrics/', {'reason': 'withdrawal'}).status_code, 200)

    def test_schedule_and_session_constraints(self):
        self.assertEqual(self.client.post(f'/api/sections/{self.section.pk}/start/', {'schedule': self.schedule.pk}).status_code, 400)
        services.close_session(self.teacher, self.session.pk)
        self.assertEqual(self.client.post(f'/api/sections/{self.section.pk}/start/', {'schedule': self.schedule.pk}).status_code, 400)
        self.assertEqual(self.client.post(f'/api/sections/{self.section.pk}/start/', {'schedule': 999}).status_code, 404)

    def test_malformed_actions_are_validation_errors(self):
        self.assertEqual(self.client.post(self.path('manual'), {'student': 'garbage', 'reason': 'x'}).status_code, 400)
        self.assertEqual(self.client.post(self.path('confirm'), {'job': 'garbage'}).status_code, 400)
        self.assertEqual(self.client.post(f'/api/sections/{self.section.pk}/start/', {'schedule': 'garbage'}).status_code, 400)
        self.assertEqual(self.client.get('/api/sessions/?section=garbage').status_code, 400)

    def test_real_cookie_login_csrf_and_staff_role(self):
        client = APIClient(enforce_csrf_checks=True)
        self.assertEqual(client.post('/api/auth/login/', {'username':'teacher','password':'password'}).status_code, 403)
        token = client.get('/api/auth/csrf/').data['csrfToken']
        self.assertEqual(client.post('/api/auth/login/', {'username':'teacher','password':'password'}, HTTP_X_CSRFTOKEN=token).status_code, 200)
        self.assertEqual(client.get('/api/auth/me/').data['role'], 'teacher')
        self.assertEqual(client.post(self.path('close'), {}).status_code, 403)
        new_token = client.get('/api/auth/csrf/').data['csrfToken']
        self.assertEqual(client.post(self.path('close'), {}, HTTP_X_CSRFTOKEN=new_token).status_code, 200)

    def test_bulk_enrollment_permissions_and_liveness_failure(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.post(f'/api/students/{self.student.pk}/bulk-enroll/', {}).status_code, 403)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.post(f'/api/students/{self.student.pk}/bulk-enroll/', {}).status_code, 404)
        self.client.force_authenticate(self.teacher)
        with patch('attendance.enrollment.extract_embedding', side_effect=RecognitionRejected('liveness_failed')):
            response = self.client.post(f'/api/students/{self.student.pk}/bulk-enroll/', {'consent':'true', 'photos':SimpleUploadedFile('f.jpg',jpeg())}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(BiometricSample.objects.count(), 0)
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        self.assertTrue(AuditLog.objects.filter(event='enrollment.rejected').exists())

    def test_withdrawal_blocks_pending_confirmation_and_private_headers(self):
        job = self.review()
        self.student.consent_at = None
        self.student.save()
        self.assertEqual(self.client.post(self.path('confirm'), {'job':str(job.pk)}).status_code, 400)
        response = self.client.get(f'/api/sessions/{self.session.pk}/')
        self.assertEqual(response['Cache-Control'], 'no-store, private')
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_oversized_and_corrupt_frame_uploads_fail_closed(self):
        token = self.client.post(self.path('capture-grant'), {}).data['token']
        bad = self.client.post(self.path('capture'), {'token':token,'frame':SimpleUploadedFile('f.jpg', b'not-an-image')}, format='multipart')
        self.assertEqual(bad.status_code, 400)
        oversized = self.client.post(self.path('capture'), {'token':token,'frame':SimpleUploadedFile('f.jpg', b'x'*(5*1024*1024+1))}, format='multipart')
        self.assertEqual(oversized.status_code, 413)
        self.assertEqual(RecognitionJob.objects.count(), 0)
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_job_visibility_admin_oversight_requester_bound_confirmation(self):
        job = self.review()
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f'/api/jobs/{job.pk}/').status_code, 404)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.get(f'/api/jobs/{job.pk}/').status_code, 200)
        self.assertEqual(self.client.post(self.path('confirm'), {'job':str(job.pk)}).status_code, 400)
